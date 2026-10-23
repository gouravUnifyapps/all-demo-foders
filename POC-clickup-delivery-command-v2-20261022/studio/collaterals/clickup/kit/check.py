# kit/check.py: the consistency gate. Run after `capture` (before recording) and again in every review.
#   python3 kit/check.py            -> review/check.md, exit 1 if anything must be fixed
# It checks, mechanically:
#   1. every number the narration says is on that scene's screen (or in facts.json)
#   2. facts.json: wherever a fact's phrase appears with a number - in narration, on any screen, in any output -
#      the number matches the fact (catches "Two samples are late" on one screen vs "4 samples late" on another)
#   3. raw data on screen: [] {} null undefined NaN {{template}}
#   4. sentences that claim an action ("sends", "updates", "approving runs...") with nothing cued on screen
#   5. internal notes in customer-facing outputs (PDF/PPTX/SRT text and speaker notes)
# What it cannot judge - whether an AI panel contradicts a record, whether a claim matches the buttons on screen -
# is the reviewer's job: review/check.md lists each scene's narration next to its screen text for that read.
import json, os, re, glob, subprocess, sys

brief = json.load(open("brief.json")); story = json.load(open("story.json"))
facts = json.load(open("facts.json")) if os.path.exists("facts.json") else []
narr = json.load(open("build/narration.json")) if os.path.exists("build/narration.json") else {}
scenes = [(a, s) for a in story["acts"] for s in a["scenes"]]
screen = {s["id"]: open(f"build/text/{s['id']}.txt").read() for _, s in scenes if os.path.exists(f"build/text/{s['id']}.txt")}

UNITS = {w: i for i, w in enumerate("zero one two three four five six seven eight nine ten eleven twelve thirteen fourteen fifteen sixteen seventeen eighteen nineteen".split())}
TENS = {w: 10 * i for i, w in enumerate("_ _ twenty thirty forty fifty sixty seventy eighty ninety".split()) if w != "_"}
ORD = {"first": 1, "second": 2, "third": 3, "fourth": 4, "fifth": 5, "sixth": 6, "seventh": 7, "eighth": 8, "ninth": 9,
       "tenth": 10, "eleventh": 11, "twelfth": 12, "thirteenth": 13, "fourteenth": 14, "fifteenth": 15, "sixteenth": 16,
       "seventeenth": 17, "eighteenth": 18, "nineteenth": 19, "twentieth": 20, "thirtieth": 30}
NUMW = set(UNITS) | set(TENS) | set(ORD) | {"hundred", "thousand", "million", "and"}

def spoken_numbers(text):
    """Numbers in a sentence, whether written as digits or words: 'one hundred and eighteen' -> 118."""
    out = [float(x.replace(",", "")) for x in re.findall(r"(?<![\w.])\d[\d,]*(?:\.\d+)?", text)]
    words = re.findall(r"[a-z]+|[^a-z\s]", text.lower()); i = 0
    while i < len(words):
        if words[i] in NUMW and words[i] != "and":
            total, cur, j = 0, 0, i
            while j < len(words) and words[j] in NUMW:
                w = words[j]
                if w in UNITS: cur += UNITS[w]
                elif w in TENS: cur += TENS[w]
                elif w in ORD: cur += ORD[w]
                elif w == "hundred": cur = max(cur, 1) * 100
                elif w in ("thousand", "million"): total += max(cur, 1) * (1000 if w == "thousand" else 10**6); cur = 0
                j += 1
            if words[j - 1] == "and": j -= 1
            out.append(float(total + cur)); i = j
        else: i += 1
    return out

def screen_numbers(text):
    vals = set()
    for x in re.findall(r"\d[\d,]*(?:\.\d+)?", text):
        v = float(x.replace(",", "")); vals |= {v, float(round(v)), float(int(v))}
    return vals

IDIOM = {0.0, 1.0}  # "one place", "one click", "one inbox": too common to check
fact_vals = {float(f["value"]) for f in facts if isinstance(f.get("value"), (int, float))}
must, verify, info = [], [], []

# 1. numbers in the narration vs the screen it plays over
for a, s in scenes:
    if s["id"] not in screen: verify.append(f"{s['id']}: no screen text yet (run capture)"); continue
    on = screen_numbers(screen[s["id"]])
    for v in spoken_numbers(s["narration"]):
        if v in IDIOM or v in on or v in fact_vals: continue
        must.append(f"{s['id']}: the voice says **{v:g}**, which is not on this screen. Narration: \"{s['narration'][:140]}…\"")

# 2. facts: same thing, same number, everywhere
texts = {f"screen {k}": v for k, v in screen.items()}
texts.update({f"narration {s['id']}": s["narration"] for _, s in scenes})
outs = {}
for f in glob.glob("out/*.pdf"):
    try: outs[f] = subprocess.run(["pdftotext", "-layout", f, "-"], capture_output=True, text=True).stdout
    except FileNotFoundError: info.append("pdftotext not installed: PDFs not scanned")
for f in glob.glob("out/*.pptx"):
    from pptx import Presentation
    p = Presentation(f); chunks = []
    for i, sl in enumerate(p.slides, 1):
        for sh in sl.shapes:
            if sh.has_text_frame: chunks.append(sh.text_frame.text)
        if sl.has_notes_slide: chunks.append(f"[slide {i} speaker notes] " + sl.notes_slide.notes_text_frame.text)
    outs[f] = "\n".join(chunks)
for f in glob.glob("out/*.srt"): outs[f] = open(f).read()
texts.update(outs)
for fct in facts:
    for phrase in fct.get("phrases", []):
        pat = re.compile(r"((?:[\w.,£$€%-]+\s+){0,3})" + re.escape(phrase.lower()))
        for where, txt in texts.items():
            for m in pat.finditer(txt.lower()):
                nums = spoken_numbers(m.group(1))
                if nums and nums[-1] != float(fct["value"]):
                    must.append(f"{where}: \"…{m.group(0).strip()}…\" says {nums[-1]:g}, but fact `{fct['id']}` is {fct['value']} ({fct.get('source', 'facts.json')})")

# 3. raw data leaking onto screens
RAW = re.compile(r"(?:^|\s)(\[\]|\{\}|null|undefined|NaN|\[object Object\]|\{\{[^}]*\}\}|Invalid Date)(?=\s|$)")
for k, txt in screen.items():
    for m in RAW.finditer(txt): must.append(f"screen {k}: raw value `{m.group(1)}` is visible on screen")

# 4. told, not shown
CLAIM = re.compile(r"\b(moves?|updates?|sends?|sent|approv(?:e|es|ing)|exports?|logs?|runs?|syncs?|notif(?:y|ies)|creates?|generates?|books?|submits?)\b", re.I)
for a, s in scenes:
    sents = [x for x in re.split(r"(?<=[.!?])\s+", s["narration"]) if x]
    cued = [x["at"] for x in s.get("actions", []) if x.get("at")]
    for x in sents:
        if CLAIM.search(x) and not any(c in x for c in cued):
            verify.append(f"{s['id']}: told, not shown? \"{x}\" (nothing on screen is cued in this sentence; show it, or check the screen has the button/state that makes it true)")

# 5. internal notes in what the customer will receive
INTERNAL = re.compile(r"\b(error|bug|fix(?:ed|es)?|todo|tbd|fixme|uat|staging|preview|localhost|sandbox|model connection|not shown|workaround|hack|internal only|confirm (?:the wording|with)|placeholder|lorem|dummy)\b", re.I)
extra = [t.lower() for t in brief.get("internal_terms", [])]
mail = re.compile(r"[\w.+-]+@[\w-]+\.[\w.]+")
for f, txt in outs.items():
    for ln in txt.splitlines():
        hit = INTERNAL.search(ln) or any(t in ln.lower() for t in extra) or (mail.search(ln) and brief.get("share") != "internal")
        if hit: must.append(f"{os.path.basename(f)}: internal note in a customer file: \"{ln.strip()[:160]}\"")

os.makedirs("review", exist_ok=True)
with open("review/check.md", "w") as r:
    r.write(f"# Consistency check\n\n**{len(must)} must fix · {len(verify)} to verify**\n\n## Must fix\n\n")
    r.write("\n".join(f"- {x}" for x in must) or "- none"); r.write("\n\n## Verify\n\n")
    r.write("\n".join(f"- {x}" for x in verify) or "- none"); r.write("\n\n")
    if info: r.write("## Notes\n\n" + "\n".join(f"- {x}" for x in info) + "\n\n")
    r.write("## Read as the viewer (do this yourself, every version)\n\nFor each scene: does anything on this screen contradict the voice, "
            "an earlier screen, or another panel on the same screen (an AI summary vs the record it summarises; a status vs a date; "
            "a claim like \"sends it\" vs the buttons that exist)? Would someone who runs this business believe every sentence?\n\n")
    for a, s in scenes:
        r.write(f"### {s['id']} · {s.get('title', '')}\n\n**Voice:** {s['narration']}\n\n**Screen text:**\n\n~~~\n{screen.get(s['id'], '(not captured)')[:2500]}\n~~~\n\n")
print(f"review/check.md: {len(must)} must fix, {len(verify)} to verify")
sys.exit(1 if must else 0)
