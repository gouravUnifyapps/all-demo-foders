# kit/deck.py: the PPT, built from the UnifyApps template. Every slide is a copy of one of the template's
# own slides (so master, background, footer, fonts and spacing carry over) with its text and pictures swapped.
#   python3 kit/deck.py --inspect      list every template slide's text slots (T#) and pictures (P#)
#   python3 kit/deck.py                build out/<Slug>_Deck.pptx from deck.json
# deck.json: {"slides": [
#   {"from": 1, "text": {"T0": "EYEBROW", "T1": ["Lead: ", "italic olive part", "."], "Merchandiser": "..."},
#    "images": {"P0": "lucide:camera", "P9": "build/shots/a2.png"}, "notes": "speaker notes"},
#   {"image": "build/swimlane_with_deck.png", "from": 1, "text": {"T0": "...", "T1": [...]}} ]}
# An image value can be {"path": "...", "y": 0.4} to choose which band of a tall screenshot fills a wide frame.
# Text keys are a slot id (T#) or a unique sample text. A list fills the runs in order (keeps the italic olive run).
# "lucide:<name>" renders that line icon in the colour of the icon it replaces. Screenshots are cropped to fill.
import sys, json, copy, io, os, re
from pptx import Presentation
from pptx.util import Emu
from PIL import Image

TPL = "brand/template.pptx"
def slots(slide):
    T, P = [], []
    for sh in sorted(slide.shapes, key=lambda s: (s.top or 0, s.left or 0)):
        if sh.shape_type == 13: P.append(sh)
        elif sh.has_text_frame and sh.text_frame.text.strip(): T.append(sh)
    return T, P

if "--inspect" in sys.argv:
    prs = Presentation(TPL)
    for i, s in enumerate(prs.slides, 1):
        T, P = slots(s); print(f"\n=== template slide {i}")
        for k, sh in enumerate(T):
            runs = [r.text for p in sh.text_frame.paragraphs for r in p.runs]
            print(f"  T{k:<3} {'|'.join(runs)[:90]!r}")
        for k, sh in enumerate(P):
            print(f"  P{k:<3} {Emu(sh.width).inches:.2f}x{Emu(sh.height).inches:.2f}in at ({Emu(sh.left).inches:.2f},{Emu(sh.top).inches:.2f})")
    sys.exit()

brief = json.load(open("brief.json")); spec = json.load(open("deck.json"))
prs = Presentation(TPL); originals = list(prs.slides)

def duplicate(src):
    dst = prs.slides.add_slide(src.slide_layout)
    for ph in list(dst.placeholders): ph._element.getparent().remove(ph._element)
    rid = {}
    for r in src.part.rels.values():
        if "image" in r.reltype: rid[r.rId] = dst.part.relate_to(r._target, r.reltype)
    for el in src.shapes._spTree.iterchildren():
        if el.tag.endswith(("}sp", "}pic", "}grpSp", "}cxnSp", "}graphicFrame")):
            x = copy.deepcopy(el); xml = None
            for b in x.iter():
                for k, v in list(b.attrib.items()):
                    if k.endswith("}embed") and v in rid: b.set(k, rid[v])
            dst.shapes._spTree.append(x)
    return dst

def set_text(sh, val):
    runs = [r for p in sh.text_frame.paragraphs for r in p.runs]
    vals = val if isinstance(val, list) else [val]
    for i, r in enumerate(runs): r.text = vals[i] if i < len(vals) else ""
    if len(vals) > len(runs): runs[-1].text += "".join(vals[len(runs):])

def colour_of(pic):
    import numpy as np
    a = np.asarray(Image.open(io.BytesIO(pic.image.blob)).convert("RGBA")).reshape(-1, 4)
    a = a[a[:, 3] > 200][:, :3]
    if not len(a): return "#847347"
    v, c = np.unique(a, axis=0, return_counts=True); return "#%02x%02x%02x" % tuple(v[c.argmax()])

def icon_png(name, colour, size=192):
    try:
        import resvg_py  # no system library needed (works on a Mac without Homebrew)
    except ImportError:
        resvg_py = None
        import cairosvg
    f = f"node_modules/lucide-static/icons/{name}.svg"
    if not os.path.exists(f):
        near = [x[:-4] for x in os.listdir(os.path.dirname(f)) if name.split("-")[0] in x][:12]
        raise SystemExit(f"no icon {name!r}; similar: {near}")
    svg = open(f).read()
    svg = svg.replace("currentColor", colour).replace('stroke-width="2"', 'stroke-width="1.75"')
    if resvg_py:
        return io.BytesIO(bytes(resvg_py.svg_to_bytes(svg_string=svg, width=size, height=size)))
    return io.BytesIO(cairosvg.svg2png(bytestring=svg.encode(), output_width=size, output_height=size))

def fill_crop(path, w, h, y=0.0):
    im = Image.open(path).convert("RGB"); r = w / h
    if im.width / im.height > r: nw = int(im.height * r); im = im.crop(((im.width - nw) // 2, 0, (im.width - nw) // 2 + nw, im.height))
    else:  # keep a band of the screenshot; y = where the band starts, as a fraction of the spare height
        nh = int(im.width / r); top = int((im.height - nh) * y); im = im.crop((0, top, im.width, top + nh))
    b = io.BytesIO(); im.save(b, "PNG"); b.seek(0); return b

def swap_picture(slide, pic, stream):
    part, rId = slide.part.get_or_add_image_part(stream)
    pic._element.blipFill.blip.rEmbed = rId

made = []
for n, sd in enumerate(spec["slides"], 1):
    src = originals[sd.get("from", 1) - 1]; s = duplicate(src); T, P = slots(s)
    sample = {}
    for k, sh in enumerate(T): sample.setdefault(sh.text_frame.text.strip(), []).append(k)
    for key, val in sd.get("text", {}).items():
        if re.fullmatch(r"T\d+", key): idx = int(key[1:])
        elif len(sample.get(key, [])) == 1: idx = sample[key][0]
        else: raise SystemExit(f"slide {n}: text key {key!r} is missing or not unique; use its T# from --inspect")
        set_text(T[idx], val)
    for key, val in sd.get("images", {}).items():
        pic = P[int(key[1:])]
        if isinstance(val, str) and val.startswith("lucide:"): swap_picture(s, pic, icon_png(val[7:], colour_of(pic)))
        else:
            v = val if isinstance(val, dict) else {"path": val}
            swap_picture(s, pic, fill_crop(v["path"], pic.width, pic.height, v.get("y", 0.0)))
    if sd.get("image"):  # a full slide picture (e.g. the swimlane) over this slide's title and footer
        keep = set(sd.get("keep", ["T0", "T1"]))
        for k, sh in enumerate(T):
            if f"T{k}" not in keep: sh._element.getparent().remove(sh._element)
        for sh in P: sh._element.getparent().remove(sh._element)
        for sh in list(s.shapes):
            if sh.shape_type == 1 or sh.shape_type == 9 or (sh.shape_type == 17 and not sh.text_frame.text.strip()):
                if sh not in T: sh._element.getparent().remove(sh._element)  # cards, connectors, empty boxes
        # Trim the picture's empty (transparent) margins, then fit it, aspect kept, into the content box
        # between the title and the footer, so nothing runs into the footer and the text stays readable.
        im = Image.open(sd["image"]).convert("RGBA"); bb = im.getchannel("A").getbbox() or (0, 0, im.width, im.height)
        im = im.crop(bb); buf = io.BytesIO(); im.save(buf, "PNG"); buf.seek(0)
        E = 914400; L, R = int(sd.get("left_in", 0.41) * E), int(sd.get("right_in", 0.41) * E)
        top, bottom = int(sd.get("top_in", 1.2) * E), prs.slide_height - int(sd.get("bottom_in", 0.6) * E)
        bw, bh = prs.slide_width - L - R, bottom - top
        w, h = (bw, int(bw * im.height / im.width)) if bw / bh < im.width / im.height else (int(bh * im.width / im.height), bh)
        s.shapes.add_picture(buf, L + (bw - w) // 2, top + (bh - h) // 2, width=w, height=h)
        if w < 0.8 * prs.slide_width: print(f"WARN slide {n}: picture only {w / prs.slide_width:.0%} of the slide width; text may be small. Use fewer lanes or steps.")
    if sd.get("notes"): s.notes_slide.notes_text_frame.text = sd["notes"]
    made.append(s)

# Drop the template's sample slides from this OUTPUT copy only (the template file itself is never changed).
lst = prs.slides._sldIdLst
for sl in originals:
    for sid in list(lst):
        if prs.slides.get(int(sid.get("id"))) is sl: prs.part.drop_rel(sid.rId); lst.remove(sid)
os.makedirs("out", exist_ok=True); out = f"out/{brief['slug']}_Deck.pptx"; prs.save(out); print(out, len(made), "slides")
